#!/usr/bin/env python3
"""Build the versioned ImpactU routing catalog from its authoritative XLSX."""

from __future__ import annotations

import argparse
import csv
from hashlib import sha256
from io import StringIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


SCHEMA_VERSION = "impactu-type-routing-v1"
CATALOG_VERSION = "1.0.0"
EXPECTED_HEADERS = ("Fuente", "Tipo", "Tipo ImpactU", "Entidad")
EXPECTED_ENTITIES = {"works", "projects", "patents", "events"}


def clean(value: Any) -> str:
    return " ".join(str(value or "").replace("\xa0", " ").split()).strip()


def build_catalog(workbook_path: Path) -> dict[str, Any]:
    workbook_bytes = workbook_path.read_bytes()
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if "ALL" not in workbook.sheetnames:
        raise ValueError("the workbook does not contain the ALL sheet")
    worksheet = workbook["ALL"]
    rows = worksheet.iter_rows(values_only=True)
    headers = tuple(clean(value) for value in next(rows))
    if headers != EXPECTED_HEADERS:
        raise ValueError(f"unexpected ALL headers: {headers!r}")

    mappings: list[dict[str, str]] = []
    seen: dict[tuple[str, str], tuple[str, str]] = {}
    for row_number, row in enumerate(rows, start=2):
        source, native_type, impactu_type, entity = (clean(value) for value in row)
        if not any((source, native_type, impactu_type, entity)):
            continue
        if not all((source, native_type, impactu_type, entity)):
            raise ValueError(f"incomplete catalog row {row_number}")
        if entity not in EXPECTED_ENTITIES:
            raise ValueError(f"invalid entity {entity!r} at row {row_number}")
        key = (source.casefold(), native_type.casefold())
        value = (impactu_type, entity)
        if key in seen and seen[key] != value:
            raise ValueError(f"conflicting route at row {row_number}: {key!r}")
        if key in seen:
            raise ValueError(f"duplicate route at row {row_number}: {key!r}")
        seen[key] = value
        mappings.append(
            {
                "source": source,
                "type": native_type,
                "type_impactu": impactu_type,
                "entity": entity,
            }
        )
    workbook.close()
    mappings.sort(key=lambda value: (value["source"].casefold(), value["type"].casefold()))
    return {
        "schema_version": SCHEMA_VERSION,
        "catalog_version": CATALOG_VERSION,
        "source": {
            "workbook": workbook_path.name,
            "sheet": "ALL",
            "sha256": sha256(workbook_bytes).hexdigest(),
            "rows": len(mappings),
        },
        "entities": sorted(EXPECTED_ENTITIES),
        "mappings": mappings,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    catalog = build_catalog(args.workbook)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    source = catalog["source"]
    buffer = StringIO(newline="")
    for key, value in (
        ("schema_version", catalog["schema_version"]),
        ("catalog_version", catalog["catalog_version"]),
        ("workbook", source["workbook"]),
        ("sheet", source["sheet"]),
        ("sha256", source["sha256"]),
        ("rows", source["rows"]),
    ):
        buffer.write(f"# {key}={value}\n")
    writer = csv.DictWriter(
        buffer,
        fieldnames=("source", "type", "type_impactu", "entity"),
        lineterminator="\n",
    )
    writer.writeheader()
    writer.writerows(catalog["mappings"])
    args.output.write_text(buffer.getvalue(), encoding="utf-8")
    print(
        f"catalog={args.output} rows={catalog['source']['rows']} "
        f"sha256={catalog['source']['sha256']}"
    )


if __name__ == "__main__":
    main()
