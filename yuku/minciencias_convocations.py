#!/usr/bin/env python3
"""Construye el histórico consolidado de convocatorias de Minciencias.

Genera un Excel con cuatro hojas normalizadas:

* investigadores
* investigadores_convocatorias
* grupos
* grupos_convocatorias

También genera dos arreglos JSON independientes, sin envoltorio general:

* investigadores.json
* grupos.json

Solo procesa PDF con texto nativo. Los documentos escaneados que requieren OCR
se excluyen deliberadamente para no incorporar errores de reconocimiento.

Uso:
    python procesar_minciencias_pdfs.py
    python procesar_minciencias_pdfs.py /ruta/Convocatorias -o historico.xlsx

Dependencias:
    Python 3.9+, openpyxl, pdftotext/pdftohtml (poppler-utils) y Ghostscript
    para abrir temporalmente PDF nativos con restricciones de copia.
"""

from __future__ import annotations

import argparse
import gc
import hashlib
import json
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata
import xml.etree.ElementTree as ET
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULT_INPUT = SCRIPT_DIR / "Convocatorias"
DEFAULT_OUTPUT = SCRIPT_DIR / "minciencias_historico.xlsx"


@dataclass(frozen=True)
class Convocatoria:
    numero: int
    año: int

    @property
    def etiqueta(self) -> str:
        return f"{self.numero}-{self.año}"

    @property
    def orden(self) -> tuple[int, int]:
        return self.año, self.numero


@dataclass(frozen=True)
class Fuente:
    pdf: Path
    relativa: str
    sha256: str
    convocatoria: Convocatoria
    clase: str


@dataclass(frozen=True)
class InvestigadorFila:
    documento_original: str
    documento_normalizado: str
    tipo_documento: str
    categoria: str
    pagina: int
    orden_fuente: int


@dataclass(frozen=True)
class GrupoFila:
    codigo: str
    nombre: str
    lider: str
    instituciones: str
    reconocido: str
    inscrito_medicion: str
    clasificacion: str
    pagina: int
    observaciones: str = ""


# Conteos de filas del listado, no necesariamente del resumen narrativo. Algunos
# documentos oficiales contienen duplicados o difieren de su texto introductorio.
ESPERADOS_INVESTIGADORES = {
    (640, 2013): 8011,
    (693, 2014): 8255,
    (737, 2015): 10042,
    (781, 2017): 13005,
    (833, 2018): 16799,
    (894, 2021): 21094,
    (957, 2024): 25514,
}

ESPERADOS_GRUPOS = {
    (359, 2006): 1437,
    (598, 2012): 5510,
    (640, 2013): 4304,
    (693, 2014): 3962,
    (737, 2015): 4627,
    (781, 2017): 5207,
    (833, 2018): 5772,
    (894, 2021): 6160,
    (957, 2024): 6322,
}

# Límites horizontales normalizados por familia de tabla moderna. Cada límite
# separa código, nombre, líder, instituciones, reconocido, inscripción y clase.
LIMITES_GRUPOS = {
    (598, 2012): (0.23, 0.45, 0.65),
    (640, 2013): (0.19, 0.35, 0.50, 0.63, 0.71, 0.79),
    (693, 2014): (0.20, 0.33, 0.45, 0.58, 0.68, 0.79),
    (737, 2015): (0.205, 0.36, 0.50, 0.62, 0.71, 0.795),
    (781, 2017): (0.215, 0.37, 0.50, 0.63, 0.725, 0.815),
    (833, 2018): (0.215, 0.37, 0.51, 0.635, 0.73, 0.81),
    (894, 2021): (0.235, 0.39, 0.52, 0.70, 0.795, 0.85),
    (957, 2024): (0.193, 0.378, 0.499, 0.652, 0.707, 0.783),
}

CLASIFICACIONES = {
    "A1",
    "A",
    "B",
    "C",
    "D",
    "Reconocido",
    "Reconocido - Sin Clasificar",
    "Sin clasificar",
    "No clasificado",
}


def argumentos() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Consolida históricamente los PDF nativos de investigadores y "
            "grupos de Minciencias en un Excel de cuatro hojas."
        )
    )
    parser.add_argument(
        "entrada",
        nargs="?",
        type=Path,
        default=DEFAULT_INPUT,
        help=f"Carpeta raíz de PDF (por defecto: {DEFAULT_INPUT}).",
    )
    parser.add_argument(
        "-o",
        "--salida",
        type=Path,
        default=DEFAULT_OUTPUT,
        help=f"Excel de salida (por defecto: {DEFAULT_OUTPUT}).",
    )
    parser.add_argument(
        "--investigadores-json",
        type=Path,
        default=None,
        help=(
            "JSON de investigadores. Por defecto se escribe investigadores.json "
            "junto al Excel."
        ),
    )
    parser.add_argument(
        "--grupos-json",
        type=Path,
        default=None,
        help=(
            "JSON de grupos. Por defecto se escribe grupos.json junto al Excel."
        ),
    )
    return parser.parse_args()


def limpiar(texto: str) -> str:
    texto = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", texto or "")
    return re.sub(r"\s+", " ", texto).strip()


def sin_acentos(texto: str) -> str:
    return "".join(
        c
        for c in unicodedata.normalize("NFD", texto)
        if unicodedata.category(c) != "Mn"
    )


def unicos(valores: Iterable[str]) -> list[str]:
    salida: list[str] = []
    for valor in valores:
        valor = limpiar(valor)
        if valor and valor not in salida:
            salida.append(valor)
    return salida


def unir(valores: Iterable[str], separador: str = " | ") -> str:
    return separador.join(unicos(valores))


def ultimo_no_vacio(
    historia: Sequence[dict[str, object]], campo: str
) -> object:
    """Devuelve el valor disponible en la convocatoria más reciente."""
    for fila in reversed(historia):
        valor = fila.get(campo)
        if valor is not None and (not isinstance(valor, str) or valor.strip()):
            return valor
    return ""


def unir_fragmentos(fragmentos: Sequence[str]) -> str:
    """Une líneas de una celda y revierte guiones de partición tipográfica."""
    resultado = ""
    for fragmento in fragmentos:
        fragmento = limpiar(fragmento)
        if not fragmento:
            continue
        ultimo = resultado.rsplit(" ", 1)[-1] if resultado else ""
        particion = (
            resultado.endswith("-")
            and len(ultimo) > 2
            and not ultimo.startswith("-")
            and bool(re.search(r"[A-Za-zÁÉÍÓÚÜÑáéíóúüñ]-$", ultimo))
            and bool(re.match(r"[A-Za-zÁÉÍÓÚÜÑáéíóúüñ]", fragmento))
        )
        if particion:
            resultado = resultado[:-1] + fragmento
        else:
            resultado = f"{resultado} {fragmento}".strip()
    return limpiar(resultado)


def ejecutar(comando: list[str]) -> str:
    try:
        proceso = subprocess.run(
            comando,
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            encoding="utf-8",
            errors="replace",
        )
    except subprocess.CalledProcessError as exc:
        detalle = limpiar(exc.stderr or exc.stdout or "error desconocido")
        raise RuntimeError(f"Falló {comando[0]}: {detalle}") from exc
    return proceso.stdout


def comprobar_dependencias() -> None:
    faltantes = [
        programa
        for programa in ("pdfinfo", "pdftotext", "pdftohtml", "gs")
        if not shutil.which(programa)
    ]
    if faltantes:
        raise RuntimeError("Faltan dependencias: " + ", ".join(faltantes))


def sha256_archivo(ruta: Path) -> str:
    digest = hashlib.sha256()
    with ruta.open("rb") as archivo:
        for bloque in iter(lambda: archivo.read(1024 * 1024), b""):
            digest.update(bloque)
    return digest.hexdigest()


def texto_pdf(pdf: Path, paginas: tuple[int, int] | None = None) -> str:
    comando = ["pdftotext", "-layout", "-enc", "UTF-8"]
    if paginas:
        comando.extend(["-f", str(paginas[0]), "-l", str(paginas[1])])
    comando.extend([str(pdf), "-"])
    return ejecutar(comando)


def pdf_cifrado(pdf: Path) -> bool:
    info = ejecutar(["pdfinfo", str(pdf)])
    match = re.search(r"^Encrypted:\s+(.+)$", info, re.MULTILINE)
    return bool(match and not match.group(1).lower().startswith("no"))


def detectar_convocatoria(texto: str, pdf: Path) -> Convocatoria | None:
    plano = limpiar(texto[:80000])
    patrones = (
        r"Convocatoria\s+(?:N(?:o|ro|°|\.)?\s*)?(\d{3})\s+(?:de|[-–])\s*(20\d{2})",
        r"Convocatoria\s+(\d{3})\D{0,20}(20\d{2})",
    )
    for patron in patrones:
        match = re.search(patron, plano, flags=re.IGNORECASE)
        if match:
            return Convocatoria(int(match.group(1)), int(match.group(2)))

    nombre = sin_acentos(str(pdf)).lower()
    conocidas = {
        "359": 2006,
        "482": 2008,
        "509": 2010,
        "542": 2011,
        "598": 2012,
        "640": 2013,
        "693": 2014,
        "737": 2015,
        "781": 2017,
        "833": 2018,
        "894": 2021,
        "957": 2024,
    }
    for numero, año in conocidas.items():
        if re.search(rf"(?<!\d){numero}(?!\d)", nombre):
            return Convocatoria(int(numero), año)
    return None


def normalizar_tipo_documento(tipo: str) -> str:
    clave = sin_acentos(limpiar(tipo)).upper()
    if "CIUDADANIA" in clave:
        return "Cédula de ciudadanía"
    if "EXTRANJER" in clave:
        return "Cédula de extranjería"
    if "PASAPORTE" in clave:
        return "Pasaporte"
    if "TARJETA" in clave:
        return "Tarjeta de identidad"
    if "REGISTRO" in clave:
        return "Registro civil"
    return limpiar(tipo).title()


def normalizar_documento(documento: str) -> str:
    documento = unicodedata.normalize("NFKC", limpiar(documento)).upper()
    return re.sub(r"[^0-9A-Z]", "", documento)


def normalizar_categoria_investigador(categoria: str) -> str:
    clave = sin_acentos(categoria).lower()
    if "emerito" in clave:
        return "Investigador Emérito (IE)"
    if "senior" in clave:
        return "Investigador Sénior (IS)"
    if "asociado" in clave:
        return "Investigador Asociado (I)"
    if "junior" in clave:
        return "Investigador Junior (IJ)"
    return limpiar(categoria)


def normalizar_si_no(valor: str) -> str:
    clave = sin_acentos(limpiar(valor)).upper().replace("\uf020", "")
    if clave in {"SI", "SÍ", "X", "✓", "√", "U", "Ü"} or valor.strip() in {
        "ü",
        "✓",
        "√",
    }:
        return "Sí"
    if clave == "NO":
        return "No"
    return limpiar(valor)


def normalizar_clasificacion(valor: str) -> tuple[str, str]:
    original = limpiar(valor)
    marcador = ""
    match_marcador = re.search(r"(\(\*+\))\s*$", original)
    if match_marcador:
        marcador = match_marcador.group(1)
        original = limpiar(original[: match_marcador.start()])
    original = re.sub(r"Reco-\s*nocido", "Reconocido", original, flags=re.I)
    clave = sin_acentos(original).lower()
    equivalencias = {
        "reconocido": "Reconocido",
        "reconocido - sin clasificar": "Reconocido - Sin Clasificar",
        "sin clasificar": "Sin clasificar",
        "no clasificado": "No clasificado",
    }
    if clave in equivalencias:
        return equivalencias[clave], marcador
    if re.fullmatch(r"a1|a|b|c|d", clave):
        return clave.upper(), marcador
    return limpiar(original), marcador


def es_clasificacion(valor: str) -> bool:
    normalizada, _ = normalizar_clasificacion(valor)
    return normalizada in CLASIFICACIONES


def clasificar_fuente(pdf: Path, raiz: Path, hash_pdf: str) -> Fuente | None:
    muestra = texto_pdf(pdf, (1, min(3, int(re.search(r"Pages:\s+(\d+)", ejecutar(["pdfinfo", str(pdf)])).group(1)))))
    convocatoria = detectar_convocatoria(muestra, pdf)
    if convocatoria is None:
        return None
    if not limpiar(muestra):
        clase = "ocr_excluido"
    elif re.search(
        r"Investigador\s+(?:Em[eé]rito|S[eé]nior|Asociado|Junior)\s*\((?:IE|IS|I|IJ)\)",
        muestra,
        flags=re.I,
    ):
        clase = "investigadores"
    elif re.search(r"COL\s*\d{4,7}", muestra, flags=re.I) or convocatoria == Convocatoria(359, 2006):
        clase = "grupos"
    else:
        clase = "no_reconocido"
    return Fuente(
        pdf=pdf,
        relativa=str(pdf.relative_to(raiz)),
        sha256=hash_pdf,
        convocatoria=convocatoria,
        clase=clase,
    )


PATRON_CATEGORIA = re.compile(
    r"(Investigador\s+(?:Em[eé]rito|S[eé]nior|Asociado|Junior)"
    r"\s*\((?:IE|IS|I|IJ)\))(?:\s*\(\*+\))?\s*$",
    flags=re.IGNORECASE,
)


def extraer_investigadores(fuente: Fuente) -> list[InvestigadorFila]:
    texto = texto_pdf(fuente.pdf)
    filas: list[InvestigadorFila] = []
    orden = 0
    for pagina, contenido in enumerate(texto.split("\f"), start=1):
        for linea in contenido.splitlines():
            match = PATRON_CATEGORIA.search(linea)
            if not match:
                continue
            prefijo = linea[: match.start()].strip()
            partes = [p.strip() for p in re.split(r"\s{2,}", prefijo) if p.strip()]
            if not partes or not re.search(r"\d", partes[-1]):
                continue
            documento = partes[-1]
            tipo = ""
            if len(partes) >= 2 and re.search(
                r"C[EÉ]DULA|PASAPORTE|TARJETA|REGISTRO|PERMISO",
                partes[-2],
                flags=re.I,
            ):
                tipo = normalizar_tipo_documento(partes[-2])
            normalizado = normalizar_documento(documento)
            if not normalizado:
                continue
            orden += 1
            filas.append(
                InvestigadorFila(
                    documento_original=limpiar(documento),
                    documento_normalizado=normalizado,
                    tipo_documento=tipo,
                    categoria=normalizar_categoria_investigador(match.group(1)),
                    pagina=pagina,
                    orden_fuente=orden,
                )
            )
    validar_investigadores_fuente(fuente, filas)
    return filas


def validar_investigadores_fuente(
    fuente: Fuente, filas: Sequence[InvestigadorFila]
) -> None:
    clave = (fuente.convocatoria.numero, fuente.convocatoria.año)
    esperado = ESPERADOS_INVESTIGADORES.get(clave)
    if esperado is None:
        raise RuntimeError(f"No existe adaptador de investigadores para {clave}")
    if len(filas) != esperado:
        raise RuntimeError(
            f"{fuente.relativa}: se esperaban {format(esperado, ',')} filas "
            f"y se extrajeron {format(len(filas), ',')}"
        )
    permitidas = {
        "Investigador Emérito (IE)",
        "Investigador Sénior (IS)",
        "Investigador Asociado (I)",
        "Investigador Junior (IJ)",
    }
    invalidas = [fila for fila in filas if fila.categoria not in permitidas]
    if invalidas:
        raise RuntimeError(f"{fuente.relativa}: categorías de investigador inválidas")


def preparar_pdf_xml(fuente: Fuente, temporal: Path) -> Path:
    pdf_usable = fuente.pdf
    if pdf_cifrado(fuente.pdf):
        pdf_usable = temporal / f"{fuente.sha256[:12]}-sin-restricciones.pdf"
        ejecutar(
            [
                "gs",
                "-q",
                "-dNOPAUSE",
                "-dBATCH",
                "-sDEVICE=pdfwrite",
                f"-sOutputFile={pdf_usable}",
                str(fuente.pdf),
            ]
        )
    xml = temporal / f"{fuente.sha256[:12]}.xml"
    ejecutar(
        [
            "pdftohtml",
            "-xml",
            "-hidden",
            "-i",
            "-q",
            str(pdf_usable),
            str(xml),
        ]
    )
    if not xml.exists():
        raise RuntimeError(f"No se generó XML para {fuente.relativa}")
    return xml


def elementos_pagina(pagina: ET.Element) -> list[tuple[float, float, str]]:
    ancho = float(pagina.attrib["width"])
    alto = float(pagina.attrib["height"])
    elementos: list[tuple[float, float, str]] = []
    for elemento in pagina.findall("text"):
        texto = limpiar("".join(elemento.itertext()))
        if not texto:
            continue
        izquierda = float(elemento.attrib["left"])
        anchura = float(elemento.attrib["width"])
        centro_x = (izquierda + anchura / 2) / ancho
        arriba = float(elemento.attrib["top"]) / alto
        elementos.append((centro_x, arriba, texto))
    return elementos


def extraer_grupos_xml(fuente: Fuente, xml: Path) -> list[GrupoFila]:
    clave = (fuente.convocatoria.numero, fuente.convocatoria.año)
    limites = LIMITES_GRUPOS[clave]
    es_2012 = clave == (598, 2012)
    filas: list[GrupoFila] = []
    raiz = ET.parse(xml).getroot()

    for pagina in raiz.findall("page"):
        numero_pagina = int(pagina.attrib["number"])
        elementos = elementos_pagina(pagina)
        inicios = [
            indice
            for indice, (x, _, texto) in enumerate(elementos)
            if x < limites[0] and re.fullmatch(r"COL\d{4,7}", texto, flags=re.I)
        ]
        for posicion, inicio in enumerate(inicios):
            fin = inicios[posicion + 1] if posicion + 1 < len(inicios) else len(elementos)
            celdas: list[list[str]] = [[] for _ in range(len(limites) + 1)]

            for x, arriba, texto in elementos[inicio:fin]:
                if (
                    es_2012
                    and x < limites[0]
                    and texto != elementos[inicio][2]
                ):
                    # El último registro de la última página va seguido de una
                    # nota narrativa en la zona del código. Esta marca termina
                    # la tabla y evita incorporar el pie al nombre del grupo.
                    break
                columna = 0
                while columna < len(limites) and x >= limites[columna]:
                    columna += 1
                celdas[columna].append(texto)
                if not es_2012 and columna == 6 and es_clasificacion(texto):
                    break

            valores = [unir_fragmentos(celda) for celda in celdas]
            codigo = re.search(
                r"COL\d{7}", re.sub(r"\s+", "", valores[0]), flags=re.I
            )
            if not codigo:
                continue

            if es_2012:
                nombre, lider, instituciones = valores[1:4]
                reconocido, inscrito, clasificacion = "Sí", "", ""
                observaciones = "La fuente corresponde al listado de grupos reconocidos."
            else:
                nombre, lider, instituciones = valores[1:4]
                reconocido = normalizar_si_no(valores[4])
                inscrito = normalizar_si_no(valores[5])
                clasificacion, marcador = normalizar_clasificacion(valores[6])
                observaciones = (
                    f"La clasificación contiene el marcador fuente {marcador}."
                    if marcador
                    else ""
                )
            filas.append(
                GrupoFila(
                    codigo=codigo.group(0).upper(),
                    nombre=nombre,
                    lider=lider,
                    instituciones=instituciones,
                    reconocido=reconocido,
                    inscrito_medicion=inscrito,
                    clasificacion=clasificacion,
                    pagina=numero_pagina,
                    observaciones=observaciones,
                )
            )
    return filas


def extraer_grupos_2006(fuente: Fuente, temporal: Path) -> list[GrupoFila]:
    tsv = temporal / f"{fuente.sha256[:12]}.tsv"
    ejecutar(["pdftotext", "-tsv", "-enc", "UTF-8", str(fuente.pdf), str(tsv)])
    palabras: dict[int, list[tuple[float, float, float, str]]] = defaultdict(list)
    with tsv.open(encoding="utf-8") as archivo:
        encabezado = next(archivo, None)
        if not encabezado:
            raise RuntimeError(f"TSV vacío para {fuente.relativa}")
        for linea in archivo:
            # No se usa csv.DictReader: nombres como '"GRUPO...' contienen una
            # comilla literal que el lector CSV interpretaría erróneamente como
            # un campo multilínea y contaminaría la celda con filas del TSV.
            campos = linea.rstrip("\r\n").split("\t", 11)
            if len(campos) != 12 or campos[0] != "5":
                continue
            try:
                palabras[int(campos[1])].append(
                    (
                        float(campos[7]),
                        float(campos[6]),
                        float(campos[8]),
                        campos[11],
                    )
                )
            except (TypeError, ValueError):
                continue

    resultado: list[GrupoFila] = []
    for pagina, elementos in palabras.items():
        anclas = sorted(
            [
                palabra
                for palabra in elementos
                if palabra[1] < 130
                and re.fullmatch(r"COL\d{7}", palabra[3], flags=re.I)
            ]
        )
        for indice, ancla in enumerate(anclas):
            inicio_y = ancla[0] - 0.5
            fin_y = anclas[indice + 1][0] - 0.5 if indice + 1 < len(anclas) else 700
            celdas: list[list[tuple[float, float, str]]] = [[] for _ in range(5)]
            for arriba, izquierda, ancho, texto in elementos:
                if not inicio_y <= arriba < fin_y:
                    continue
                centro = izquierda + ancho / 2
                columna = (
                    0
                    if centro < 130
                    else 1
                    if centro < 260
                    else 2
                    if centro < 470
                    else 3
                    if centro < 580
                    else 4
                )
                celdas[columna].append((arriba, izquierda, texto))
            valores = [
                limpiar(" ".join(item[2] for item in sorted(celda))) for celda in celdas
            ]
            categoria = ""
            match_categoria = re.search(r"\b([ABC])\b", valores[4])
            if match_categoria:
                categoria = match_categoria.group(1)
            resultado.append(
                GrupoFila(
                    codigo=ancla[3].upper(),
                    nombre=valores[1],
                    lider=valores[3],
                    instituciones=valores[2],
                    reconocido="Sí",
                    inscrito_medicion="",
                    clasificacion=categoria,
                    pagina=pagina,
                    observaciones="La fuente corresponde a la medición de grupos reconocidos.",
                )
            )
    return resultado


def texto_desde_palabras(
    palabras: Sequence[tuple[float, float, str]],
) -> str:
    lineas: list[list[object]] = []
    for arriba, izquierda, texto in sorted(palabras):
        for linea in lineas:
            if abs(arriba - float(linea[0])) < 1.5:
                linea[1].append((izquierda, texto))  # type: ignore[union-attr]
                break
        else:
            lineas.append([arriba, [(izquierda, texto)]])
    fragmentos = [
        " ".join(texto for _, texto in sorted(linea[1]))  # type: ignore[arg-type]
        for linea in lineas
    ]
    return unir_fragmentos(fragmentos)


def extraer_grupos_2012(fuente: Fuente, temporal: Path) -> list[GrupoFila]:
    """Extrae la tabla de 2012 por palabras y bandas verticales.

    En algunas filas Poppler combina nombre y líder en un solo bloque XML. El
    TSV por palabra conserva sus coordenadas y permite separarlos sin inferir.
    """
    tsv = temporal / f"{fuente.sha256[:12]}-2012.tsv"
    ejecutar(["pdftotext", "-tsv", "-enc", "UTF-8", str(fuente.pdf), str(tsv)])
    palabras: dict[int, list[tuple[float, float, float, str]]] = defaultdict(list)
    with tsv.open(encoding="utf-8") as archivo:
        next(archivo, None)
        for linea in archivo:
            campos = linea.rstrip("\r\n").split("\t", 11)
            if len(campos) != 12 or campos[0] != "5":
                continue
            try:
                palabras[int(campos[1])].append(
                    (float(campos[7]), float(campos[6]), float(campos[8]), campos[11])
                )
            except (TypeError, ValueError):
                continue

    resultado: list[GrupoFila] = []
    for pagina, elementos in palabras.items():
        anclas = sorted(
            [
                palabra
                for palabra in elementos
                if palabra[1] < 140
                and re.fullmatch(r"COL\d{7}", palabra[3], flags=re.I)
            ]
        )
        posiciones_y = [ancla[0] for ancla in anclas]
        for indice, ancla in enumerate(anclas):
            inicio_y = (
                (posiciones_y[indice - 1] + posiciones_y[indice]) / 2
                if indice
                else posiciones_y[indice]
                - (posiciones_y[indice + 1] - posiciones_y[indice]) / 2
                if len(posiciones_y) > 1
                else posiciones_y[indice] - 20
            )
            fin_y = (
                (posiciones_y[indice] + posiciones_y[indice + 1]) / 2
                if indice + 1 < len(posiciones_y)
                else posiciones_y[indice]
                + (posiciones_y[indice] - posiciones_y[indice - 1]) / 2
                if len(posiciones_y) > 1
                else posiciones_y[indice] + 20
            )
            celdas: list[list[tuple[float, float, str]]] = [[] for _ in range(4)]
            for arriba, izquierda, ancho, texto in elementos:
                if not inicio_y <= arriba < fin_y:
                    continue
                centro = izquierda + ancho / 2
                columna = 0 if centro < 138 else 1 if centro < 285 else 2 if centro < 395 else 3
                celdas[columna].append((arriba, izquierda, texto))
            valores = [texto_desde_palabras(celda) for celda in celdas]
            resultado.append(
                GrupoFila(
                    codigo=ancla[3].upper(),
                    nombre=valores[1],
                    lider=valores[2],
                    instituciones=valores[3],
                    reconocido="Sí",
                    inscrito_medicion="",
                    clasificacion="",
                    pagina=pagina,
                    observaciones="La fuente corresponde al listado de grupos reconocidos.",
                )
            )
    return resultado


def extraer_grupos(fuente: Fuente, temporal: Path) -> list[GrupoFila]:
    clave = (fuente.convocatoria.numero, fuente.convocatoria.año)
    if clave == (359, 2006):
        filas = extraer_grupos_2006(fuente, temporal)
    elif clave == (598, 2012):
        filas = extraer_grupos_2012(fuente, temporal)
    elif clave in LIMITES_GRUPOS:
        filas = extraer_grupos_xml(fuente, preparar_pdf_xml(fuente, temporal))
    else:
        raise RuntimeError(f"No existe adaptador de grupos para {clave}")
    validar_grupos_fuente(fuente, filas)
    return filas


def validar_grupos_fuente(fuente: Fuente, filas: Sequence[GrupoFila]) -> None:
    clave = (fuente.convocatoria.numero, fuente.convocatoria.año)
    esperado = ESPERADOS_GRUPOS.get(clave)
    if esperado is None:
        raise RuntimeError(f"No existe conteo esperado de grupos para {clave}")
    if len(filas) != esperado:
        raise RuntimeError(
            f"{fuente.relativa}: se esperaban {format(esperado, ',')} grupos "
            f"y se extrajeron {format(len(filas), ',')}"
        )
    codigos = [fila.codigo for fila in filas]
    if len(codigos) != len(set(codigos)):
        raise RuntimeError(f"{fuente.relativa}: códigos de grupo duplicados")
    if any(not re.fullmatch(r"COL\d{7}", codigo) for codigo in codigos):
        raise RuntimeError(f"{fuente.relativa}: código de grupo inválido")
    invalidas = [
        fila.clasificacion
        for fila in filas
        if fila.clasificacion and fila.clasificacion not in CLASIFICACIONES
    ]
    if invalidas:
        raise RuntimeError(
            f"{fuente.relativa}: clasificaciones inválidas: {sorted(set(invalidas))}"
        )
    if any(not fila.nombre for fila in filas):
        raise RuntimeError(f"{fuente.relativa}: existen grupos sin nombre")
    contaminadas = [
        fila.codigo
        for fila in filas
        if "###LINE###" in " ".join(
            [fila.nombre, fila.lider, fila.instituciones]
        )
    ]
    if contaminadas:
        raise RuntimeError(
            f"{fuente.relativa}: celdas contaminadas por estructura PDF: "
            f"{contaminadas[:5]}"
        )


def id_investigador(clave_identidad: str) -> str:
    resumen = hashlib.sha256(clave_identidad.encode("utf-8")).hexdigest()[:16].upper()
    return f"INV-{resumen}"


def observaciones_campos_vacios(fila: GrupoFila) -> list[str]:
    vacios = []
    if not fila.lider:
        vacios.append("líder")
    if not fila.instituciones:
        vacios.append("instituciones")
    if not vacios:
        return []
    return ["Campos vacíos en la fuente: " + ", ".join(vacios) + "."]


def consolidar_investigadores(
    extracciones: Sequence[tuple[Fuente, Sequence[InvestigadorFila]]],
) -> tuple[list[dict[str, object]], list[dict[str, object]]]:
    detalles_agrupados: dict[tuple[str, str], list[tuple[Fuente, InvestigadorFila]]] = defaultdict(list)
    ceros_por_fuente: Counter[tuple[str, str]] = Counter()
    tipos_por_documento: dict[str, set[str]] = defaultdict(set)

    for _, filas in extracciones:
        for fila in filas:
            if fila.tipo_documento:
                tipos_por_documento[fila.documento_normalizado].add(
                    fila.tipo_documento
                )

    for fuente, filas in extracciones:
        for fila in filas:
            if set(fila.documento_normalizado) == {"0"}:
                contador_clave = (fuente.sha256, fila.documento_normalizado)
                ceros_por_fuente[contador_clave] += 1
                identidad = (
                    f"CERO: {fuente.convocatoria.etiqueta}: {fila.documento_normalizado}: "
                    f"{ceros_por_fuente[contador_clave]}"
                )
            else:
                # El número normalizado es la única clave transversal presente
                # en todas las convocatorias. Los tipos históricos se conservan
                # y sus inconsistencias se documentan, sin romper la trayectoria.
                identidad = f"DOC: {fila.documento_normalizado}"
            detalles_agrupados[(identidad, fuente.convocatoria.etiqueta)].append(
                (fuente, fila)
            )

    detalles: list[dict[str, object]] = []
    identidad_detalles: dict[str, list[dict[str, object]]] = defaultdict(list)
    for (identidad, _), ocurrencias in detalles_agrupados.items():
        fuente = ocurrencias[0][0]
        filas = [item[1] for item in ocurrencias]
        categorias = unicos(fila.categoria for fila in filas)
        tipos = unicos(fila.tipo_documento for fila in filas)
        originales = unicos(fila.documento_original for fila in filas)
        paginas = sorted(set(fila.pagina for fila in filas))
        observaciones: list[str] = []
        if len(filas) > 1:
            observaciones.append(
                f"La fuente contiene {len(filas)} filas para esta identidad y convocatoria."
            )
        if len(categorias) > 1:
            observaciones.append("La fuente asigna categorías incompatibles al mismo documento.")
        if len(tipos) > 1:
            observaciones.append("La fuente registra más de un tipo de documento.")
        if set(filas[0].documento_normalizado) == {"0"}:
            observaciones.append(
                "Documento no identificable; no se vinculó con otras convocatorias."
            )
        if len(tipos_por_documento[filas[0].documento_normalizado]) > 1:
            observaciones.append(
                "El número aparece con tipos de documento diferentes entre "
                "las fuentes; la vinculación se realizó por número normalizado."
            )
        detalle = {
            "investigador_id": id_investigador(identidad),
            "convocatoria": fuente.convocatoria.etiqueta,
            "numero_convocatoria": fuente.convocatoria.numero,
            "año": fuente.convocatoria.año,
            "tipo_documento_original": unir(tipos),
            "documento_original": unir(originales),
            "documento_normalizado": filas[0].documento_normalizado,
            "categoria": unir(categorias),
            "archivo_fuente": fuente.relativa,
            "pagina_fuente": "; ".join(str(pagina) for pagina in paginas),
            "metodo_extraccion": "texto_nativo",
            "estado_validacion": (
                "validado_con_advertencia" if observaciones else "validado"
            ),
            "observaciones": " ".join(observaciones),
        }
        detalles.append(detalle)
        identidad_detalles[identidad].append(detalle)

    maestros: list[dict[str, object]] = []
    for identidad, historia in identidad_detalles.items():
        historia.sort(key=lambda fila: (int(fila["año"]), int(fila["numero_convocatoria"])))
        tipos = unicos(str(fila["tipo_documento_original"]) for fila in historia)
        categorias = unicos(str(fila["categoria"]) for fila in historia)
        convocatorias = [str(fila["convocatoria"]) for fila in historia]
        advertencias = unicos(str(fila["observaciones"]) for fila in historia)
        maestros.append(
            {
                "investigador_id": historia[0]["investigador_id"],
                "tipo_documento_consolidado": ultimo_no_vacio(
                    historia, "tipo_documento_original"
                ),
                "tipos_documento_historicos": unir(tipos),
                "documento_original_preferido": ultimo_no_vacio(
                    historia, "documento_original"
                ),
                "documento_normalizado": historia[0]["documento_normalizado"],
                "categorias_historicas": unir(categorias),
                "convocatorias": "; ".join(convocatorias),
                "numero_convocatorias": len(convocatorias),
                "primera_convocatoria": convocatorias[0],
                "ultima_convocatoria": convocatorias[-1],
                "estado_validacion": (
                    "validado_con_advertencia" if advertencias else "validado"
                ),
                "observaciones": " ".join(advertencias),
            }
        )

    maestros.sort(key=lambda fila: (str(fila["documento_normalizado"]), str(fila["investigador_id"])))
    detalles.sort(key=lambda fila: (int(fila["año"]), str(fila["documento_normalizado"]), str(fila["investigador_id"])))
    return maestros, detalles


def consolidar_grupos(
    extracciones: Sequence[tuple[Fuente, Sequence[GrupoFila]]],
) -> tuple[list[dict[str, object]], list[dict[str, object]]]:
    detalles: list[dict[str, object]] = []
    por_grupo: dict[str, list[dict[str, object]]] = defaultdict(list)
    vistos: set[tuple[str, str]] = set()

    for fuente, filas in extracciones:
        for fila in filas:
            clave = (fila.codigo, fuente.convocatoria.etiqueta)
            if clave in vistos:
                raise RuntimeError(f"Grupo/convocatoria duplicado: {clave}")
            vistos.add(clave)
            observaciones = unicos(
                [fila.observaciones, *observaciones_campos_vacios(fila)]
            )
            detalle = {
                "grupo_id": f"GRP-{fila.codigo}",
                "codigo_grupo": fila.codigo,
                "convocatoria": fuente.convocatoria.etiqueta,
                "numero_convocatoria": fuente.convocatoria.numero,
                "año": fuente.convocatoria.año,
                "nombre_grupo": fila.nombre,
                "lider_grupo": fila.lider,
                "instituciones_avaladoras": fila.instituciones,
                "grupo_reconocido": fila.reconocido,
                "inscrito_medicion": fila.inscrito_medicion,
                "clasificacion": fila.clasificacion,
                "archivo_fuente": fuente.relativa,
                "pagina_fuente": fila.pagina,
                "metodo_extraccion": "texto_nativo",
                "estado_validacion": (
                    "validado_con_advertencia" if observaciones else "validado"
                ),
                "observaciones": " ".join(observaciones),
            }
            detalles.append(detalle)
            por_grupo[fila.codigo].append(detalle)

    maestros: list[dict[str, object]] = []
    for codigo, historia in por_grupo.items():
        historia.sort(key=lambda fila: (int(fila["año"]), int(fila["numero_convocatoria"])))
        nombres = unicos(str(fila["nombre_grupo"]) for fila in historia)
        lideres = unicos(str(fila["lider_grupo"]) for fila in historia)
        instituciones = unicos(str(fila["instituciones_avaladoras"]) for fila in historia)
        clasificaciones = unicos(str(fila["clasificacion"]) for fila in historia)
        convocatorias = [str(fila["convocatoria"]) for fila in historia]
        advertencias = unicos(str(fila["observaciones"]) for fila in historia)
        maestros.append(
            {
                "grupo_id": f"GRP-{codigo}",
                "codigo_grupo": codigo,
                "nombre_preferido": ultimo_no_vacio(historia, "nombre_grupo"),
                "nombres_historicos": unir(nombres),
                "lider_mas_reciente": ultimo_no_vacio(historia, "lider_grupo"),
                "lideres_historicos": unir(lideres),
                "instituciones_historicas": unir(instituciones),
                "clasificaciones_historicas": unir(clasificaciones),
                "convocatorias": "; ".join(convocatorias),
                "numero_convocatorias": len(convocatorias),
                "primera_convocatoria": convocatorias[0],
                "ultima_convocatoria": convocatorias[-1],
                "estado_validacion": (
                    "validado_con_advertencia" if advertencias else "validado"
                ),
                "observaciones": " ".join(advertencias),
            }
        )

    maestros.sort(key=lambda fila: str(fila["codigo_grupo"]))
    detalles.sort(key=lambda fila: (int(fila["año"]), str(fila["codigo_grupo"])))
    return maestros, detalles


ENCABEZADOS_INVESTIGADORES = [
    "investigador_id",
    "tipo_documento_consolidado",
    "tipos_documento_historicos",
    "documento_original_preferido",
    "documento_normalizado",
    "categorias_historicas",
    "convocatorias",
    "numero_convocatorias",
    "primera_convocatoria",
    "ultima_convocatoria",
    "estado_validacion",
    "observaciones",
]

ENCABEZADOS_INVESTIGADORES_CONV = [
    "investigador_id",
    "convocatoria",
    "numero_convocatoria",
    "año",
    "tipo_documento_original",
    "documento_original",
    "documento_normalizado",
    "categoria",
    "archivo_fuente",
    "pagina_fuente",
    "metodo_extraccion",
    "estado_validacion",
    "observaciones",
]

ENCABEZADOS_GRUPOS = [
    "grupo_id",
    "codigo_grupo",
    "nombre_preferido",
    "nombres_historicos",
    "lider_mas_reciente",
    "lideres_historicos",
    "instituciones_historicas",
    "clasificaciones_historicas",
    "convocatorias",
    "numero_convocatorias",
    "primera_convocatoria",
    "ultima_convocatoria",
    "estado_validacion",
    "observaciones",
]

ENCABEZADOS_GRUPOS_CONV = [
    "grupo_id",
    "codigo_grupo",
    "convocatoria",
    "numero_convocatoria",
    "año",
    "nombre_grupo",
    "lider_grupo",
    "instituciones_avaladoras",
    "grupo_reconocido",
    "inscrito_medicion",
    "clasificacion",
    "archivo_fuente",
    "pagina_fuente",
    "metodo_extraccion",
    "estado_validacion",
    "observaciones",
]


def valor_json(valor: object) -> object:
    """Convierte los campos vacíos del Excel en null sin alterar los demás."""
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return None
    return valor


def valores_compuestos(valor: object) -> list[str]:
    """Separa valores consolidados preservando su orden de aparición."""
    if valor is None:
        return []
    return unicos(parte.strip() for parte in str(valor).split(" | "))


def resumen_json(historia: Sequence[dict[str, object]]) -> dict[str, object]:
    primera = historia[0]
    ultima = historia[-1]
    return {
        "primera_convocatoria": primera["convocatoria"],
        "ultima_convocatoria": ultima["convocatoria"],
        "año_primera_convocatoria": primera["año"],
        "año_ultima_convocatoria": ultima["año"],
        "numero_convocatorias": len(historia),
    }


def escribir_arreglo_json(salida: Path, registros: Iterable[dict[str, object]]) -> int:
    """Escribe un arreglo JSON de forma incremental y reemplaza atómicamente."""
    salida.parent.mkdir(parents=True, exist_ok=True)
    temporal = salida.with_name(f".{salida.name}.tmp")
    cantidad = 0
    try:
        with temporal.open("w", encoding="utf-8", newline="\n") as archivo:
            archivo.write("[\n")
            for registro in registros:
                if cantidad:
                    archivo.write(",\n")
                serializado = json.dumps(
                    registro,
                    ensure_ascii=False,
                    indent=2,
                    allow_nan=False,
                )
                archivo.write("  " + serializado.replace("\n", "\n  "))
                cantidad += 1
            archivo.write("\n]\n")
        temporal.replace(salida)
    except Exception:
        temporal.unlink(missing_ok=True)
        raise
    return cantidad


def escribir_json_investigadores(
    salida: Path,
    investigadores: Sequence[dict[str, object]],
    detalles: Sequence[dict[str, object]],
) -> int:
    historias: dict[str, list[dict[str, object]]] = defaultdict(list)
    for detalle in detalles:
        historias[str(detalle["investigador_id"])].append(detalle)

    def registros() -> Iterable[dict[str, object]]:
        for maestro in investigadores:
            investigador_id = str(maestro["investigador_id"])
            historia = historias[investigador_id]
            historia.sort(
                key=lambda fila: (
                    int(fila["año"]),
                    int(fila["numero_convocatoria"]),
                )
            )
            tipos = unicos(
                tipo
                for fila in historia
                for tipo in valores_compuestos(fila["tipo_documento_original"])
            )
            categorias = [
                {
                    "convocatoria": fila["convocatoria"],
                    "año": fila["año"],
                    "categoria": valor_json(fila["categoria"]),
                }
                for fila in historia
            ]
            yield {
                "investigador_id": investigador_id,
                "documento_original": valor_json(
                    maestro["documento_original_preferido"]
                ),
                "documento_normalizado": valor_json(
                    maestro["documento_normalizado"]
                ),
                "tipo_consolidado": valor_json(
                    maestro["tipo_documento_consolidado"]
                ),
                "tipos_historicos": tipos,
                "categorias_historicas": categorias,
                "resumen_historico": resumen_json(historia),
            }

    return escribir_arreglo_json(salida, registros())


def escribir_json_grupos(
    salida: Path,
    grupos: Sequence[dict[str, object]],
    detalles: Sequence[dict[str, object]],
) -> int:
    historias: dict[str, list[dict[str, object]]] = defaultdict(list)
    for detalle in detalles:
        historias[str(detalle["grupo_id"])].append(detalle)

    def historico(
        historia: Sequence[dict[str, object]], campo_origen: str, campo_json: str
    ) -> list[dict[str, object]]:
        return [
            {
                "convocatoria": fila["convocatoria"],
                "año": fila["año"],
                campo_json: valor_json(fila[campo_origen]),
            }
            for fila in historia
        ]

    def registros() -> Iterable[dict[str, object]]:
        for maestro in grupos:
            grupo_id = str(maestro["grupo_id"])
            historia = historias[grupo_id]
            historia.sort(
                key=lambda fila: (
                    int(fila["año"]),
                    int(fila["numero_convocatoria"]),
                )
            )
            yield {
                "grupo_id": grupo_id,
                "codigo_grupo": maestro["codigo_grupo"],
                "nombre_preferido": valor_json(maestro["nombre_preferido"]),
                "nombres_historicos": historico(
                    historia, "nombre_grupo", "nombre"
                ),
                "lider_mas_reciente": valor_json(maestro["lider_mas_reciente"]),
                "lideres_historicos": historico(
                    historia, "lider_grupo", "lider"
                ),
                "instituciones_historicas": historico(
                    historia, "instituciones_avaladoras", "institucion"
                ),
                "clasificaciones_historicas": historico(
                    historia, "clasificacion", "clasificacion"
                ),
                "resumen_historico": resumen_json(historia),
            }

    return escribir_arreglo_json(salida, registros())


def verificar_json(
    investigadores_json: Path,
    grupos_json: Path,
    cantidades: dict[str, int],
) -> None:
    with investigadores_json.open(encoding="utf-8") as archivo:
        investigadores = json.load(archivo)
    if not isinstance(investigadores, list):
        raise RuntimeError("investigadores.json: la raíz no es un arreglo")
    if len(investigadores) != cantidades["investigadores"]:
        raise RuntimeError("investigadores.json: cantidad de registros incorrecta")
    ids = set()
    relaciones = 0
    claves_i = {
        "investigador_id",
        "documento_original",
        "documento_normalizado",
        "tipo_consolidado",
        "tipos_historicos",
        "categorias_historicas",
        "resumen_historico",
    }
    for registro in investigadores:
        if set(registro) != claves_i:
            raise RuntimeError("investigadores.json: estructura inesperada")
        investigador_id = registro["investigador_id"]
        if investigador_id in ids:
            raise RuntimeError("investigadores.json: investigador_id duplicado")
        ids.add(investigador_id)
        historia = registro["categorias_historicas"]
        convocatorias = [fila["convocatoria"] for fila in historia]
        if len(convocatorias) != len(set(convocatorias)):
            raise RuntimeError("investigadores.json: convocatoria duplicada")
        resumen = registro["resumen_historico"]
        if not historia or resumen["numero_convocatorias"] != len(historia):
            raise RuntimeError("investigadores.json: resumen histórico inconsistente")
        if (
            resumen["primera_convocatoria"] != historia[0]["convocatoria"]
            or resumen["ultima_convocatoria"] != historia[-1]["convocatoria"]
            or resumen["año_primera_convocatoria"] != historia[0]["año"]
            or resumen["año_ultima_convocatoria"] != historia[-1]["año"]
        ):
            raise RuntimeError("investigadores.json: límites históricos inconsistentes")
        relaciones += len(historia)
    if relaciones != cantidades["investigadores_convocatorias"]:
        raise RuntimeError("investigadores.json: relaciones históricas incompletas")
    del investigadores
    gc.collect()

    with grupos_json.open(encoding="utf-8") as archivo:
        grupos = json.load(archivo)
    if not isinstance(grupos, list):
        raise RuntimeError("grupos.json: la raíz no es un arreglo")
    if len(grupos) != cantidades["grupos"]:
        raise RuntimeError("grupos.json: cantidad de registros incorrecta")
    ids = set()
    relaciones = 0
    claves_g = {
        "grupo_id",
        "codigo_grupo",
        "nombre_preferido",
        "nombres_historicos",
        "lider_mas_reciente",
        "lideres_historicos",
        "instituciones_historicas",
        "clasificaciones_historicas",
        "resumen_historico",
    }
    campos_historicos = (
        "nombres_historicos",
        "lideres_historicos",
        "instituciones_historicas",
        "clasificaciones_historicas",
    )
    for registro in grupos:
        if set(registro) != claves_g:
            raise RuntimeError("grupos.json: estructura inesperada")
        grupo_id = registro["grupo_id"]
        if grupo_id in ids:
            raise RuntimeError("grupos.json: grupo_id duplicado")
        ids.add(grupo_id)
        historia = registro["nombres_historicos"]
        convocatorias = [fila["convocatoria"] for fila in historia]
        if len(convocatorias) != len(set(convocatorias)):
            raise RuntimeError("grupos.json: convocatoria duplicada")
        for campo in campos_historicos[1:]:
            otra_historia = registro[campo]
            if [fila["convocatoria"] for fila in otra_historia] != convocatorias:
                raise RuntimeError("grupos.json: históricos desalineados")
        resumen = registro["resumen_historico"]
        if not historia or resumen["numero_convocatorias"] != len(historia):
            raise RuntimeError("grupos.json: resumen histórico inconsistente")
        if (
            resumen["primera_convocatoria"] != historia[0]["convocatoria"]
            or resumen["ultima_convocatoria"] != historia[-1]["convocatoria"]
            or resumen["año_primera_convocatoria"] != historia[0]["año"]
            or resumen["año_ultima_convocatoria"] != historia[-1]["año"]
        ):
            raise RuntimeError("grupos.json: límites históricos inconsistentes")
        nombres = [fila["nombre"] for fila in historia if fila["nombre"] is not None]
        if nombres and registro["nombre_preferido"] != nombres[-1]:
            raise RuntimeError("grupos.json: nombre preferido no es el más reciente")
        lideres = [
            fila["lider"]
            for fila in registro["lideres_historicos"]
            if fila["lider"] is not None
        ]
        if lideres and registro["lider_mas_reciente"] != lideres[-1]:
            raise RuntimeError("grupos.json: líder reciente inconsistente")
        relaciones += len(historia)
    if relaciones != cantidades["grupos_convocatorias"]:
        raise RuntimeError("grupos.json: relaciones históricas incompletas")


def configurar_hoja(hoja, encabezados: Sequence[str]) -> None:
    relleno = PatternFill("solid", fgColor="1F4E78")
    fuente = Font(color="FFFFFF", bold=True)
    alineacion = Alignment(horizontal="center", vertical="center", wrap_text=True)
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = f"A1: {get_column_letter(len(encabezados))}1"
    celdas = []
    for encabezado in encabezados:
        celda = WriteOnlyCell(hoja, value=encabezado)
        celda.fill = relleno
        celda.font = fuente
        celda.alignment = alineacion
        celdas.append(celda)
    hoja.append(celdas)


def escribir_hoja(hoja, encabezados: Sequence[str], filas: Sequence[dict[str, object]]) -> None:
    anchos = []
    for encabezado in encabezados:
        if encabezado in {"observaciones", "instituciones_avaladoras", "instituciones_historicas"}:
            ancho = 60
        elif encabezado in {"nombre_grupo", "nombre_preferido", "nombres_historicos"}:
            ancho = 50
        elif "archivo_fuente" in encabezado or "historicos" in encabezado:
            ancho = 45
        elif "convocatoria" in encabezado or "categoria" in encabezado or "clasificacion" in encabezado:
            ancho = 30
        else:
            ancho = max(14, min(28, len(encabezado) + 3))
        anchos.append(ancho)
    for indice, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(indice)].width = ancho
    configurar_hoja(hoja, encabezados)
    columnas_texto = {
        indice + 1
        for indice, encabezado in enumerate(encabezados)
        if "documento" in encabezado or encabezado.endswith("_id") or encabezado == "codigo_grupo"
    }
    for fila in filas:
        valores = [fila.get(encabezado, "") for encabezado in encabezados]
        celdas = []
        for columna, valor in enumerate(valores, start=1):
            celda = WriteOnlyCell(hoja, value=valor)
            celda.alignment = Alignment(vertical="top", wrap_text=True)
            if columna in columnas_texto:
                celda.number_format = "@"
            celdas.append(celda)
        hoja.append(celdas)


def escribir_excel(
    salida: Path,
    investigadores: Sequence[dict[str, object]],
    investigadores_conv: Sequence[dict[str, object]],
    grupos: Sequence[dict[str, object]],
    grupos_conv: Sequence[dict[str, object]],
) -> None:
    salida.parent.mkdir(parents=True, exist_ok=True)
    libro = Workbook(write_only=True)
    hoja = libro.create_sheet("investigadores")
    escribir_hoja(hoja, ENCABEZADOS_INVESTIGADORES, investigadores)
    escribir_hoja(
        libro.create_sheet("investigadores_convocatorias"),
        ENCABEZADOS_INVESTIGADORES_CONV,
        investigadores_conv,
    )
    escribir_hoja(
        libro.create_sheet("grupos"), ENCABEZADOS_GRUPOS, grupos
    )
    escribir_hoja(
        libro.create_sheet("grupos_convocatorias"),
        ENCABEZADOS_GRUPOS_CONV,
        grupos_conv,
    )
    temporal = salida.with_name(f".{salida.name}.tmp.xlsx")
    libro.save(temporal)
    temporal.replace(salida)


def verificar_excel(
    salida: Path,
    cantidades_esperadas: dict[str, int],
) -> None:
    libro = load_workbook(salida, read_only=True, data_only=True)
    esperadas = [
        "investigadores",
        "investigadores_convocatorias",
        "grupos",
        "grupos_convocatorias",
    ]
    if libro.sheetnames != esperadas:
        raise RuntimeError(f"Hojas inesperadas: {libro.sheetnames}")
    conjuntos = {
        "investigadores": "investigador_id",
        "investigadores_convocatorias": ("investigador_id", "convocatoria"),
        "grupos": "grupo_id",
        "grupos_convocatorias": ("grupo_id", "convocatoria"),
    }
    for nombre, clave in conjuntos.items():
        hoja = libro[nombre]
        iterador = hoja.iter_rows(values_only=True)
        encabezados = list(next(iterador))
        posiciones = {encabezado: indice for indice, encabezado in enumerate(encabezados)}
        claves_excel = set()
        cantidad = 0
        for fila in iterador:
            cantidad += 1
            if isinstance(clave, tuple):
                valor_clave = tuple(fila[posiciones[campo]] for campo in clave)
            else:
                valor_clave = fila[posiciones[clave]]
            if valor_clave in claves_excel:
                raise RuntimeError(f"{nombre}: la clave no es única")
            claves_excel.add(valor_clave)
        if cantidad != cantidades_esperadas[nombre]:
            raise RuntimeError(
                f"{nombre}: {cantidad} filas en Excel, "
                f"{cantidades_esperadas[nombre]} esperadas"
            )


def inventariar(raiz: Path) -> tuple[list[Fuente], list[str], list[str]]:
    fuentes: list[Fuente] = []
    excluidos: list[str] = []
    duplicados: list[str] = []
    hashes_vistos: dict[str, str] = {}
    for pdf in sorted(raiz.rglob("*.pdf")):
        hash_pdf = sha256_archivo(pdf)
        relativa = str(pdf.relative_to(raiz))
        if hash_pdf in hashes_vistos:
            duplicados.append(f"{relativa} = {hashes_vistos[hash_pdf]}")
            continue
        hashes_vistos[hash_pdf] = relativa
        fuente = clasificar_fuente(pdf, raiz, hash_pdf)
        if fuente is None:
            excluidos.append(f"{relativa}: convocatoria no identificada")
        elif fuente.clase == "ocr_excluido":
            excluidos.append(f"{relativa}: escaneado sin texto nativo")
        elif fuente.clase == "no_reconocido":
            excluidos.append(f"{relativa}: formato no reconocido")
        else:
            fuentes.append(fuente)
    fuentes.sort(key=lambda f: (*f.convocatoria.orden, f.clase, f.relativa))
    return fuentes, excluidos, duplicados


def main() -> int:
    args = argumentos()
    raiz = args.entrada.resolve()
    salida = args.salida.resolve()
    investigadores_json = (
        args.investigadores_json.resolve()
        if args.investigadores_json is not None
        else salida.parent / "investigadores.json"
    )
    grupos_json = (
        args.grupos_json.resolve()
        if args.grupos_json is not None
        else salida.parent / "grupos.json"
    )
    if not raiz.is_dir():
        print(f"ERROR: no existe la carpeta {raiz}", file=sys.stderr)
        return 2
    try:
        if len({salida, investigadores_json, grupos_json}) != 3:
            raise RuntimeError("Las tres rutas de salida deben ser diferentes")
        comprobar_dependencias()
        fuentes, excluidos, duplicados = inventariar(raiz)
        if not fuentes:
            raise RuntimeError("No se encontraron fuentes nativas procesables")

        extracciones_i: list[tuple[Fuente, Sequence[InvestigadorFila]]] = []
        extracciones_g: list[tuple[Fuente, Sequence[GrupoFila]]] = []
        with tempfile.TemporaryDirectory(prefix="minciencias_historico_") as temporal:
            carpeta_temporal = Path(temporal)
            for fuente in fuentes:
                if fuente.clase == "investigadores":
                    filas_i = extraer_investigadores(fuente)
                    extracciones_i.append((fuente, filas_i))
                    print(
                        f"[investigadores {fuente.convocatoria.etiqueta}] "
                        f"{format(len(filas_i), ',')} filas — {fuente.relativa}"
                    )
                elif fuente.clase == "grupos":
                    filas_g = extraer_grupos(fuente, carpeta_temporal)
                    extracciones_g.append((fuente, filas_g))
                    print(
                        f"[grupos {fuente.convocatoria.etiqueta}] "
                        f"{format(len(filas_g), ',')} filas — {fuente.relativa}"
                    )

        investigadores, investigadores_conv = consolidar_investigadores(extracciones_i)
        grupos, grupos_conv = consolidar_grupos(extracciones_g)
        cantidades = {
            "investigadores": len(investigadores),
            "investigadores_convocatorias": len(investigadores_conv),
            "grupos": len(grupos),
            "grupos_convocatorias": len(grupos_conv),
        }
        escribir_excel(
            salida, investigadores, investigadores_conv, grupos, grupos_conv
        )
        cantidad_i_json = escribir_json_investigadores(
            investigadores_json, investigadores, investigadores_conv
        )
        cantidad_g_json = escribir_json_grupos(grupos_json, grupos, grupos_conv)
        if cantidad_i_json != cantidades["investigadores"]:
            raise RuntimeError("No se escribieron todos los investigadores en JSON")
        if cantidad_g_json != cantidades["grupos"]:
            raise RuntimeError("No se escribieron todos los grupos en JSON")
        del (
            extracciones_i,
            extracciones_g,
            investigadores,
            investigadores_conv,
            grupos,
            grupos_conv,
        )
        gc.collect()
        verificar_excel(salida, cantidades)
        verificar_json(investigadores_json, grupos_json, cantidades)

        print("\nVerificación completada")
        print(f"Investigadores únicos: {format(cantidades['investigadores'], ',')}")
        print(
            "Relaciones investigador-convocatoria: "
            f"{format(cantidades['investigadores_convocatorias'], ',')}"
        )
        print(f"Grupos únicos: {format(cantidades['grupos'], ',')}")
        print(
            "Relaciones grupo-convocatoria: "
            f"{format(cantidades['grupos_convocatorias'], ',')}"
        )
        print(f"PDF nativos procesados: {len(fuentes)}")
        print(f"PDF excluidos: {len(excluidos)}")
        print(f"PDF duplicados omitidos: {len(duplicados)}")
        print(f"Excel: {salida}")
        print(f"JSON investigadores: {investigadores_json}")
        print(f"JSON grupos: {grupos_json}")
        for mensaje in excluidos:
            print(f"EXCLUIDO: {mensaje}")
        for mensaje in duplicados:
            print(f"DUPLICADO OMITIDO: {mensaje}")
        return 0
    except (RuntimeError, ET.ParseError, OSError, ValueError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
